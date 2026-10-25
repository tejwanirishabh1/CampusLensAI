from pathlib import Path
import re

import fitz
from pptx import Presentation
from docx import Document


def clean_text(text):
    """Clean excessive whitespace while preserving paragraph breaks."""

    text = text.replace("\r\n", "\n").replace("\r", "\n")

    lines = [
        re.sub(r"[ \t]+", " ", line).strip()
        for line in text.split("\n")
    ]

    cleaned_lines = []
    previous_blank = True

    for line in lines:
        if not line:
            if cleaned_lines and not previous_blank:
                cleaned_lines.append("")

            previous_blank = True
        else:
            cleaned_lines.append(line)
            previous_blank = False

    return "\n".join(cleaned_lines).strip()


def extract_text_from_document(file_path):
    """Extract readable content from supported document formats."""

    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".pptx":
        return extract_text_from_pptx(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    raise ValueError(
        f"Unsupported file format: {extension}. "
        "Please use PDF, PPTX, or DOCX."
    )


def extract_text_from_pdf(pdf_path):
    """Extract text from each PDF page."""

    extracted_pages = []

    with fitz.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            raw_text = page.get_text("text")
            text = clean_text(raw_text)

            if text:
                extracted_pages.append({
                    "page": page_number,
                    "text": text,
                    "display_text": format_plain_text(text),
                    "source_type": "page",
                })

    return extracted_pages


def extract_text_from_pptx(pptx_path):
    """Extract slide content and distinguish likely slide titles."""

    presentation = Presentation(str(pptx_path))
    extracted_slides = []

    for slide_number, slide in enumerate(
        presentation.slides, start=1
    ):
        content_blocks = []
        display_blocks = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                paragraphs = []

                for paragraph in shape.text_frame.paragraphs:
                    paragraph_text = paragraph.text.strip()

                    if paragraph_text:
                        paragraphs.append(paragraph_text)

                if not paragraphs:
                    continue

                shape_text = "\n".join(paragraphs)
                content_blocks.append(shape_text)

                # Use PowerPoint's title placeholder when available.
                if shape.is_placeholder:
                    try:
                        is_title = (
                            shape.placeholder_format.type
                            in (1, 3)
                        )
                    except (AttributeError, ValueError):
                        is_title = False
                else:
                    is_title = False

                if is_title:
                    display_blocks.append(
                        f"## {shape_text}"
                    )
                else:
                    display_blocks.append(
                        format_plain_text(shape_text)
                    )

            if shape.has_table:
                for row in shape.table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    if any(cells):
                        row_text = " | ".join(cells)
                        content_blocks.append(row_text)
                        display_blocks.append(row_text)

        text = clean_text("\n\n".join(content_blocks))
        display_text = "\n\n".join(
            block for block in display_blocks if block.strip()
        )

        if text:
            extracted_slides.append({
                "page": slide_number,
                "text": text,
                "display_text": display_text,
                "source_type": "slide",
            })

    return extracted_slides


def extract_text_from_docx(docx_path):
    """Extract Word paragraphs, headings, lists, and tables."""

    document = Document(str(docx_path))

    content_blocks = []
    display_blocks = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        content_blocks.append(text)

        style_name = ""
        if paragraph.style is not None:
            style_name = paragraph.style.name.lower()

        # Use the document's actual heading styles.
        if style_name.startswith("heading"):
            display_blocks.append(f"## {text}")

        elif style_name == "title":
            display_blocks.append(f"# {text}")

        elif style_name == "subtitle":
            display_blocks.append(f"### {text}")

        elif "list bullet" in style_name:
            display_blocks.append(f"- {text}")

        elif "list number" in style_name:
            display_blocks.append(f"1. {text}")

        else:
            display_blocks.append(text)

    # Extract tables and separate cells clearly.
    for table in document.tables:
        table_rows = []

        for row in table.rows:
            cells = [
                clean_text(cell.text)
                for cell in row.cells
            ]

            if any(cells):
                table_rows.append(" | ".join(cells))

        if table_rows:
            content_blocks.append("\n".join(table_rows))
            display_blocks.append(
                "**Table**\n\n" + "\n\n".join(table_rows)
            )

    combined_text = clean_text("\n\n".join(content_blocks))
    display_text = "\n\n".join(
        block for block in display_blocks if block.strip()
    )

    if not combined_text:
        return []

    return [{
        "page": 1,
        "text": combined_text,
        "display_text": display_text,
        "source_type": "document",
    }]


def format_plain_text(text):
    """Prepare extracted text for a readable Markdown preview."""

    text = clean_text(text)

    if not text:
        return ""

    lines = text.split("\n")
    formatted_lines = []

    for line in lines:
        line = line.strip()

        if not line:
            continue

        # Display common bullet markers consistently.
        if re.match(r"^[•●▪]\s*", line):
            line = re.sub(r"^[•●▪]\s*", "- ", line)

        formatted_lines.append(line)

    return "\n\n".join(formatted_lines)